#!/usr/bin/env python3
"""Merge 9.11 / 9.14 / 9.15 sheets into 2026-09-10.xlsx."""
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

import openpyxl

ROOT = Path("/Users/cusky/Desktop/kakobuy")
XLSX = ROOT / "2026-09-10.xlsx"
SOURCES = [
    Path("/Users/cusky/Desktop/商品清单/2026-09-11.xlsx"),
    Path("/Users/cusky/Desktop/商品清单/黄总9.14新增商品.xlsx"),
    Path("/Users/cusky/Desktop/商品表单/9.15黄总新增商品.xlsx"),
]
UA = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
    "Accept": "*/*",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://weidian.com/",
}
BATCH = "IMP-20260916-huang"
NOW = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

META_TAGS = {
    "hot nike",
    "top trending",
    "new arrivals",
    "hot picks",
    "scarce lv",
    "streetwear",
    "football",
}
CAT_TAGS = {
    "shoes": "Shoes",
    "slides": "Shoes",
    "sneakers": "Shoes",
    "jordan 4": "Jordan 4",
    "t-shirt": "Tops",
    "shirts": "Shirts",
    "shirt": "Shirts",
    "polo": "Polo",
    "hoodies": "Hoodies",
    "hoodie": "Hoodies",
    "sweater": "Hoodies",
    "jacket": "Jacket",
    "down jacket": "Jacket",
    "vest": "Jacket",
    "shorts": "Shorts",
    "pants": "Pants",
    "jeans": "Bottoms",
    "bottoms": "Bottoms",
    "bag": "Bag",
    "backpack": "Bag",
    "accessories": "Accessories",
    "cap": "Accessories",
    "hat": "Accessories",
    "belt": "Accessories",
    "socks": "Accessories",
    "sunglasses": "Accessories",
    "bracelet": "Accessories",
    "necklace": "Accessories",
    "balaclava": "Accessories",
    "watch": "Watch",
    "electronics": "Electronics",
    "underwear": "Underwear",
    "nike tech": "Nike Tech",
    "set": "Set",
    "tracksuit": "Tracksuit",
    "suit": "Suit",
    "suirt": "Suit",
    "jersey": "Jersey",
    "dress": "Dress",
}
BRAND_FIX = {
    "ysl": "YSL",
    "adidas": "Adidas",
    "maison margiela": "Maison Margiela",
    "miumiu": "Miu Miu",
    "miu miu": "Miu Miu",
    "lv": "LV",
    "saint laurent": "Saint Laurent",
}


def item_id_from_url(url):
    if not url:
        return ""
    m = re.search(r"itemID=(\d+)", str(url), re.I)
    if m:
        return m.group(1)
    m = re.search(r"(\d{8,})", str(url))
    return m.group(1) if m else ""


def parse_tags(raw):
    parts = [p.strip() for p in str(raw or "").replace("\n", ",").split(",") if p.strip()]
    brand = ""
    cat = ""
    for part in parts:
        key = part.lower()
        if key in META_TAGS:
            continue
        if key in CAT_TAGS:
            if not cat:
                cat = CAT_TAGS[key]
            continue
        if not brand:
            brand = BRAND_FIX.get(key, part)
    if not brand:
        for part in parts:
            if part.lower() not in META_TAGS:
                brand = BRAND_FIX.get(part.lower(), part)
                break
    return brand, cat


def load_simple(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    items = []
    seen = set()
    for r in rows[1:]:
        if not r:
            continue
        title = str(r[0] or "").strip()
        url = str(r[1] or "").strip()
        tags = str(r[2] or "").strip() if len(r) > 2 else ""
        iid = item_id_from_url(url)
        if not iid or iid in seen or not title:
            continue
        seen.add(iid)
        brand, cat = parse_tags(tags)
        items.append(
            {
                "itemId": iid,
                "title": title,
                "url": url or f"https://weidian.com/item.html?itemID={iid}",
                "brand": brand,
                "cat": cat,
            }
        )
    return items


def fetch_sku(item_id):
    param = quote(json.dumps({"itemId": str(item_id)}, separators=(",", ":")))
    api = f"https://thor.weidian.com/detail/getItemSkuInfo/1.0?param={param}"
    last = ""
    for attempt in range(3):
        try:
            req = Request(api, headers=UA)
            with urlopen(req, timeout=18) as r:
                data = json.loads(r.read().decode("utf-8", "replace"))
            res = data.get("result") or {}
            raw = res.get("itemDiscountLowPrice")
            price = 0.0
            if raw not in (None, ""):
                price = float(raw) / 100.0
            return {
                "title": str(res.get("itemTitle") or "").strip(),
                "price": price,
                "ok": True,
            }
        except Exception as e:
            last = type(e).__name__
            time.sleep(0.5 * (attempt + 1))
    return {"title": "", "price": 0.0, "ok": False, "err": last}


def col_index(header):
    return {name: i + 1 for i, name in enumerate(header)}


def main():
    incoming = []
    seen = set()
    for path in SOURCES:
        chunk = load_simple(path)
        print(path.name, len(chunk), flush=True)
        for it in chunk:
            if it["itemId"] in seen:
                continue
            seen.add(it["itemId"])
            incoming.append(it)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col = col_index(header)
    id_col = col["商品id"]
    disp_col = col["展示商品标题"]

    index = {}
    for row in range(2, ws.max_row + 1):
        iid = ws.cell(row, id_col).value
        if iid:
            index[str(iid).strip()] = row

    updated = 0
    new_items = []
    for it in incoming:
        row = index.get(it["itemId"])
        if row:
            old = str(ws.cell(row, disp_col).value or "").strip()
            if old != it["title"]:
                ws.cell(row, disp_col).value = it["title"]
                updated += 1
        else:
            new_items.append(it)

    print("title updates", updated, "new rows", len(new_items), flush=True)

    sku = {}
    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = {ex.submit(fetch_sku, it["itemId"]): it["itemId"] for it in new_items}
        for i, fut in enumerate(as_completed(futs), 1):
            iid = futs[fut]
            sku[iid] = fut.result()
            if i % 50 == 0 or i == len(futs):
                print(f"sku {i}/{len(futs)}", flush=True)

    def put(row, name, value):
        if name in col:
            ws.cell(row, col[name]).value = value

    start = ws.max_row + 1
    for n, it in enumerate(new_items):
        row = start + n
        info = sku.get(it["itemId"]) or {}
        orig = info.get("title") or it["title"]
        price = info.get("price") or 0
        put(row, "id", str(2100000000000000000 + int(it["itemId"])))
        put(row, "平台", "weidian")
        put(row, "商品id", it["itemId"])
        put(row, "商品url", it["url"])
        put(row, "商品标题", orig)
        put(row, "展示商品标题", it["title"])
        put(row, "品牌", it["brand"] or None)
        put(row, "品牌分类", it["cat"] or None)
        put(row, "商品优惠价", price or None)
        put(row, "导入批次", BATCH)
        put(row, "打开商品详情页数", 0)
        put(row, "跳转数", 0)
        put(row, "CTR", 0)
        put(row, "CVR", 0)
        put(row, "状态", "1")
        put(row, "创建时间", NOW)

    wb.save(XLSX)
    sku_ok = sum(1 for v in sku.values() if v.get("ok"))
    print(
        "saved",
        XLSX.name,
        "rows",
        ws.max_row - 1,
        "sku_ok",
        sku_ok,
        "/",
        len(new_items),
        flush=True,
    )


if __name__ == "__main__":
    main()
